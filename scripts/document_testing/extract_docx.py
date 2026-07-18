#!/usr/bin/env python3
"""
Word Document (.docx) Text Extractor script using python-docx.
Usage: python extract_docx.py <path_to_docx_file>
"""

import sys
import os

def extract_docx(filepath: str) -> str:
    try:
        import docx
    except ImportError:
        return "❌ Error: 'python-docx' package is not installed. Install it with: pip install python-docx"

    if not os.path.exists(filepath):
        return f"❌ Error: File not found at '{filepath}'"

    try:
        doc = docx.Document(filepath)
        text_blocks = []
        text_blocks.append(f"📄 Document: {os.path.basename(filepath)}\n" + "="*50)

        # 1. Paragraphs
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        if paragraphs:
            text_blocks.append("=== PARAGRAPHS ===")
            text_blocks.extend(paragraphs)

        # 2. Tables
        if doc.tables:
            text_blocks.append("\n=== TABLES ===")
            for i, table in enumerate(doc.tables):
                text_blocks.append(f"\n--- Table {i+1} ---")
                for row in table.rows:
                    row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    text_blocks.append(" | ".join(row_data))

        return "\n".join(text_blocks)
    except Exception as e:
        return f"❌ Error extracting DOCX: {e}"

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python extract_docx.py <path_to_docx_file>")
        sys.exit(1)

    target_path = sys.argv[1]
    result = extract_docx(target_path)
    print(result)
