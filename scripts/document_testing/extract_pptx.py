#!/usr/bin/env python3
"""
PowerPoint Presentation (.pptx) Text Extractor script using python-pptx.
Usage: python extract_pptx.py <path_to_pptx_file>
"""

import sys
import os

def extract_pptx(filepath: str) -> str:
    try:
        import pptx
    except ImportError:
        return "❌ Error: 'python-pptx' package is not installed. Install it with: pip install python-pptx"

    if not os.path.exists(filepath):
        return f"❌ Error: File not found at '{filepath}'"

    try:
        prs = pptx.Presentation(filepath)
        text_blocks = []
        text_blocks.append(f"🖥️ Presentation: {os.path.basename(filepath)} | Slides: {len(prs.slides)}\n" + "="*50)

        for i, slide in enumerate(prs.slides):
            text_blocks.append(f"\n=== Slide {i+1} ===")
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            if slide_text:
                text_blocks.extend(slide_text)
            else:
                text_blocks.append("(No Text Found on Slide)")

        return "\n".join(text_blocks)
    except Exception as e:
        return f"❌ Error extracting PPTX: {e}"

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python extract_pptx.py <path_to_pptx_file>")
        sys.exit(1)

    target_path = sys.argv[1]
    result = extract_pptx(target_path)
    print(result)
