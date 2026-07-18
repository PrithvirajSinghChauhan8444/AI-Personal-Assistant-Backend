"""
Word Document (.docx) Parser Module.
Extracts paragraphs and tabular data using python-docx.
"""

import os

def extract_docx(filepath: str) -> str:
    """Extracts text and table structures from Word (.docx) files using python-docx."""
    try:
        import docx
    except ImportError:
        return "❌ Error: 'python-docx' package is not installed. Run: pip install python-docx"

    if not os.path.exists(filepath):
        return f"❌ Error: File not found at '{filepath}'"

    try:
        doc = docx.Document(filepath)
        text_blocks = [f"📄 Document: {os.path.basename(filepath)}\n" + "="*50]
        
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        if paragraphs:
            text_blocks.append("=== PARAGRAPHS ===")
            text_blocks.extend(paragraphs)

        if doc.tables:
            text_blocks.append("\n=== TABLES ===")
            for i, table in enumerate(doc.tables):
                text_blocks.append(f"\n--- Table {i+1} ---")
                for row in table.rows:
                    row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    text_blocks.append(" | ".join(row_data))

        return "\n".join(text_blocks)
    except Exception as e:
        return f"❌ Error extracting DOCX: {e}"
