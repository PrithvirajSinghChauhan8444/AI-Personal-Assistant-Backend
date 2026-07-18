"""
PowerPoint Presentation (.pptx) Parser Module.
Extracts slide-by-slide text using python-pptx.
"""

import os

def extract_pptx(filepath: str) -> str:
    """Extracts slide-by-slide text from PowerPoint (.pptx) files using python-pptx."""
    try:
        import pptx
    except ImportError:
        return "❌ Error: 'python-pptx' package is not installed. Run: pip install python-pptx"

    if not os.path.exists(filepath):
        return f"❌ Error: File not found at '{filepath}'"

    try:
        prs = pptx.Presentation(filepath)
        text_blocks = [f"🖥️ Presentation: {os.path.basename(filepath)} | Slides: {len(prs.slides)}\n" + "="*50]

        for i, slide in enumerate(prs.slides):
            text_blocks.append(f"\n=== Slide {i+1} ===")
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            if slide_text:
                text_blocks.extend(slide_text)
            else:
                text_blocks.append("(No Text Found on Slide)")

        return "\n".join(text_blocks)
    except Exception as e:
        return f"❌ Error extracting PPTX: {e}"
