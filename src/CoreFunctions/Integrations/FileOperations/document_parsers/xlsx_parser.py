"""
Excel Spreadsheet (.xlsx) Parser Module.
Extracts cell values sheet-by-sheet using openpyxl.
"""

import os

def extract_xlsx(filepath: str) -> str:
    """Extracts sheet-by-sheet data from Excel (.xlsx) files using openpyxl."""
    try:
        import openpyxl
    except ImportError:
        return "❌ Error: 'openpyxl' package is not installed. Run: pip install openpyxl"

    if not os.path.exists(filepath):
        return f"❌ Error: File not found at '{filepath}'"

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        text_blocks = [f"📊 Spreadsheet: {os.path.basename(filepath)} | Sheets: {len(wb.sheetnames)}\n" + "="*50]

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text_blocks.append(f"\n=== Sheet: {sheet_name} ===")
            has_data = False
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    has_data = True
                    row_data = [str(cell) if cell is not None else "" for cell in row]
                    text_blocks.append(" | ".join(row_data))
            if not has_data:
                text_blocks.append("(Empty Sheet)")

        return "\n".join(text_blocks)
    except Exception as e:
        return f"❌ Error extracting XLSX: {e}"
